#!/usr/bin/env python3
"""
OCH workbook integrity check.

Run this at the START of any session, before editing anything, and again
before committing back to the repo.

    python3 check_workbook.py OCH_Plan_Library_Template_1.xlsx

It answers one question: is this the complete, current workbook, or a
branch that's missing work done elsewhere?

Exit code 0 = clean. 1 = something is missing or regressed.

WHY THIS EXISTS
---------------
On 9/11/26 three copies of this workbook were in circulation. The one with
the newest-looking filename was OLDER on the Grand Oaks tabs and NEWER on
the Heights tabs. Neither date nor filename identified the current file.
Only content did. This script checks content.

MAINTENANCE
-----------
When a work stream completes, add a check for it in EXPECTED below. The
check is only as good as what's listed here.
"""

import sys
import openpyxl

# --- What a complete workbook must contain -------------------------------
# Add a line here whenever a work stream lands.

EXPECTED_TABS = [
    "Read Me",
    "Plan Library",
    "Community Offerings",
    "Community Standards",
    "Lot Constraints — GO Ph1A",
    "Lot Constraints — Heights",
]

# Plans that must be present in Plan Library, with their elevations.
EXPECTED_PLANS = {
    # Heights — 7 active, all rear-entry
    "Addison":    ["A", "B", "C"],
    "Carmichael": ["A", "B", "C"],
    "Crowder":    ["A", "B", "C"],
    "Holland":    ["A", "B", "C"],
    "Knox":       ["A", "B", "C"],
    "Parker":     ["A", "B", "C"],
    "Pearl":      ["A"],
    # Grand Oaks — 40' front-entry family
    "Renner":     ["A", "B", "C"],
    "Shelby":     ["D", "E", "F"],
    "Webber":     ["A", "B", "C"],
}

# Lot Constraints row counts (data rows, excluding headers).
EXPECTED_LOT_COUNTS = {
    "Lot Constraints — Heights": 48,
}

# Community Standards: (row, column, label, substring that must appear)
# Column F = The Heights at Uptown.
EXPECTED_STANDARDS = [
    (11, 6, "Heights front building line",   "10'-0\""),
    (13, 6, "Heights side setback interior",  "5'-0\""),
    (14, 6, "Heights side setback corner",   "15'-0\""),
    (15, 6, "Heights rear setback alley",    "20'-0\""),
    (36, 6, "Heights porch encroachment",     "5'-0\""),
    (43, 6, "Heights monotony intervening",   "3"),
]

# Plan Library Notes that must be present (the 40' storage matrix).
STORAGE_NOTE_PLANS = ["Renner", "Shelby", "Webber"]
STORAGE_NOTE_MARKER = "OPT STOR"


def main(path):
    problems = []
    warnings = []

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        print(f"FAIL — could not open {path}: {e}")
        return 1

    print(f"Checking: {path}\n")

    # --- Tabs ---
    print("TABS")
    present = wb.sheetnames
    for t in EXPECTED_TABS:
        ok = t in present
        print(f"  {'ok  ' if ok else 'MISS'} {t}")
        if not ok:
            problems.append(f"Missing tab: {t}")
    for t in present:
        if t not in EXPECTED_TABS:
            print(f"  ?    {t}  (unexpected — new tab, or renamed)")
            warnings.append(f"Unexpected tab: {t}")

    # --- Plans ---
    print("\nPLAN LIBRARY")
    if "Plan Library" in present:
        ws = wb["Plan Library"]
        found = {}
        notes = {}
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row[0]:
                found.setdefault(row[0], []).append(row[1])
                notes[(row[0], row[1])] = row[15] if len(row) > 15 else None
        for plan, elevs in EXPECTED_PLANS.items():
            got = found.get(plan, [])
            missing = [e for e in elevs if e not in got]
            if not got:
                print(f"  MISS {plan} — plan absent entirely")
                problems.append(f"Plan missing from library: {plan}")
            elif missing:
                print(f"  MISS {plan} — elevations missing: {missing}")
                problems.append(f"{plan} missing elevations {missing}")
            else:
                print(f"  ok   {plan} {'/'.join(elevs)}")

        # 40' storage matrix
        print("\n  40' storage matrix (Notes column)")
        for plan in STORAGE_NOTE_PLANS:
            rows = [(p, e) for (p, e) in notes if p == plan]
            have = [k for k in rows if notes[k] and STORAGE_NOTE_MARKER in str(notes[k])]
            if len(have) == len(rows) and rows:
                print(f"    ok   {plan} — {len(have)}/{len(rows)} rows")
            else:
                print(f"    MISS {plan} — {len(have)}/{len(rows)} rows carry storage note")
                problems.append(f"{plan} storage matrix incomplete ({len(have)}/{len(rows)})")
    else:
        problems.append("Plan Library tab missing — cannot check plans")

    # --- Lot constraints ---
    print("\nLOT CONSTRAINTS")
    for tab, expected in EXPECTED_LOT_COUNTS.items():
        if tab in present:
            ws = wb[tab]
            n = sum(1 for r in ws.iter_rows(min_row=4, values_only=True) if r[0])
            ok = n == expected
            print(f"  {'ok  ' if ok else 'FAIL'} {tab}: {n} lots (expected {expected})")
            if not ok:
                problems.append(f"{tab} has {n} lots, expected {expected}")
        else:
            print(f"  MISS {tab}")

    # --- Community standards ---
    print("\nCOMMUNITY STANDARDS")
    if "Community Standards" in present:
        ws = wb["Community Standards"]
        for row, col, label, marker in EXPECTED_STANDARDS:
            val = ws.cell(row=row, column=col).value
            ok = val is not None and str(marker) in str(val)
            print(f"  {'ok  ' if ok else 'MISS'} {label}: {str(val)[:45] if val else '(blank)'}")
            if not ok:
                problems.append(f"Standard missing or changed: {label} (expected to contain {marker})")
    else:
        problems.append("Community Standards tab missing")

    # --- Open flags (informational, not failures) ---
    print("\nOPEN FLAGS carried in this workbook")
    flags = 0
    for tab in present:
        ws = wb[tab]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and isinstance(cell, str):
                    up = cell.upper()
                    if "CONFLICT" in up or "DISCREPANCY" in up or "UNRESOLVED" in up:
                        print(f"  • {tab}: {cell[:90]}")
                        flags += 1
    if flags == 0:
        print("  (none)")

    # --- Verdict ---
    print("\n" + "=" * 60)
    if problems:
        print(f"RESULT: {len(problems)} PROBLEM(S) — this is NOT the complete workbook.")
        for p in problems:
            print(f"  ! {p}")
        print("\nDo not edit this file. Find the copy that has the missing work,")
        print("diff the two, and merge before doing anything else.")
    else:
        print("RESULT: CLEAN — all expected content present.")
    if warnings:
        print(f"\n{len(warnings)} warning(s):")
        for w in warnings:
            print(f"  ? {w}")
    print("=" * 60)

    return 1 if problems else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
